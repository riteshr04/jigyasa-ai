import os
import json
import tempfile
import unittest
from pathlib import Path
from io import BytesIO
from unittest.mock import MagicMock
from unittest.mock import patch

from docx import Document as WordDocument
from fastapi.testclient import TestClient
from fastapi import HTTPException, Response
from pptx import Presentation

import backend.main as backend_main
from backend.main import _extract_text, _generate_gemini_answer, _select_context


class DocumentExtractionTests(unittest.TestCase):
    def test_extracts_text_file(self):
        self.assertEqual(_extract_text("notes.txt", b"\xef\xbb\xbfCell biology\n\nCells are alive."), "Cell biology\nCells are alive.")

    def test_extracts_word_document(self):
        document = WordDocument()
        document.add_paragraph("A cell is the basic unit of life.")
        output = BytesIO()
        document.save(output)
        self.assertIn("basic unit of life", _extract_text("notes.docx", output.getvalue()))

    def test_extracts_presentation_text(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Photosynthesis"
        output = BytesIO()
        presentation.save(output)
        self.assertIn("Photosynthesis", _extract_text("lesson.pptx", output.getvalue()))

    def test_rejects_invalid_pdf_signature(self):
        with self.assertRaises(HTTPException) as context:
            _extract_text("notes.pdf", b"not a PDF")
        self.assertEqual(context.exception.status_code, 400)

    def test_rejects_empty_document(self):
        with self.assertRaises(HTTPException) as context:
            _extract_text("empty.txt", b" \n\t")
        self.assertEqual(context.exception.status_code, 422)

    def test_rejects_unsupported_format(self):
        with self.assertRaises(HTTPException) as context:
            _extract_text("lesson.mp4", b"video")
        self.assertEqual(context.exception.status_code, 415)


class StudyContextTests(unittest.TestCase):
    def test_context_ranks_matching_document_sections(self):
        text = ("Astronomy describes stars and galaxies. " * 50) + ("Mitochondria release energy inside cells. " * 50)
        context = _select_context(text, "What do mitochondria do?", limit=500)
        self.assertIn("Mitochondria", context)

    def test_gemini_requires_a_key_without_making_a_request(self):
        with patch.dict(os.environ, {"GEMINI_API_KEY": ""}):
            with self.assertRaises(HTTPException) as context:
                _generate_gemini_answer("What is a cell?", [{"name": "notes.txt", "text": "A cell is a unit of life."}])
        self.assertEqual(context.exception.status_code, 503)

    def test_gemini_request_uses_rest_camel_case_and_header_key(self):
        fake_response = MagicMock()
        fake_response.__enter__.return_value.read.return_value = json.dumps({
            "candidates": [{"content": {"parts": [{"text": "A cell is the basic unit of life."}]}}],
        }).encode()
        with (
            patch.dict(os.environ, {"GEMINI_API_KEY": "test-jigyasa-key"}),
            patch("backend.main.urlopen", return_value=fake_response) as mocked_urlopen,
        ):
            answer = _generate_gemini_answer(
                "What is a cell?",
                [{"name": "notes.txt", "text": "A cell is the basic unit of life."}],
            )

        request = mocked_urlopen.call_args.args[0]
        payload = json.loads(request.data)
        self.assertEqual(answer, "A cell is the basic unit of life.")
        self.assertIn("systemInstruction", payload)
        self.assertNotIn("system_instruction", payload)
        self.assertEqual(request.get_header("X-goog-api-key"), "test-jigyasa-key")
        self.assertNotIn("test-jigyasa-key", request.full_url)

    def test_gemini_http_400_reports_provider_reason_without_api_key(self):
        from urllib.error import HTTPError

        error_response = HTTPError(
            "https://example.invalid",
            400,
            "Bad Request",
            {},
            BytesIO(json.dumps({"error": {"message": "Invalid request payload."}}).encode()),
        )
        with (
            patch.dict(os.environ, {"GEMINI_API_KEY": "test-jigyasa-key"}),
            patch("backend.main.urlopen", side_effect=error_response),
        ):
            with self.assertRaises(HTTPException) as context:
                _generate_gemini_answer(
                    "What is a cell?",
                    [{"name": "notes.txt", "text": "A cell is the basic unit of life."}],
                )
        self.assertEqual(context.exception.status_code, 502)
        self.assertIn("Invalid request payload", context.exception.detail)
        self.assertNotIn("test-jigyasa-key", context.exception.detail)

    def test_gemini_retries_once_after_service_overload(self):
        from urllib.error import HTTPError

        fake_response = MagicMock()
        fake_response.__enter__.return_value.read.return_value = json.dumps({
            "candidates": [{"content": {"parts": [{"text": "A cell is the basic unit of life."}]}}],
        }).encode()
        overload = HTTPError("https://example.invalid", 503, "Unavailable", {}, BytesIO(b""))
        with (
            patch.dict(os.environ, {"GEMINI_API_KEY": "test-jigyasa-key"}),
            patch("backend.main.urlopen", side_effect=[overload, fake_response]) as mocked_urlopen,
            patch("backend.main.time.sleep") as mocked_sleep,
        ):
            answer = _generate_gemini_answer(
                "What is a cell?",
                [{"name": "notes.txt", "text": "A cell is the basic unit of life."}],
            )
        self.assertEqual(answer, "A cell is the basic unit of life.")
        self.assertEqual(mocked_urlopen.call_count, 2)
        mocked_sleep.assert_called_once_with(1)

    def test_gemini_reports_service_overload_after_one_retry(self):
        from urllib.error import HTTPError

        overload = HTTPError("https://example.invalid", 503, "Unavailable", {}, BytesIO(b""))
        with (
            patch.dict(os.environ, {"GEMINI_API_KEY": "test-jigyasa-key"}),
            patch("backend.main.urlopen", side_effect=overload) as mocked_urlopen,
            patch("backend.main.time.sleep"),
        ):
            with self.assertRaises(HTTPException) as context:
                _generate_gemini_answer(
                    "What is a cell?",
                    [{"name": "notes.txt", "text": "A cell is the basic unit of life."}],
                )
        self.assertEqual(context.exception.status_code, 503)
        self.assertIn("high demand", context.exception.detail)
        self.assertEqual(mocked_urlopen.call_count, 2)


class HostedSecurityTests(unittest.TestCase):
    def test_production_documents_require_a_valid_session(self):
        with (
            patch.object(backend_main, "PRODUCTION", True),
            patch.object(backend_main, "_supabase_config", return_value=("https://example.supabase.co", "anon", "service")),
            patch.dict(os.environ, {"APP_ORIGIN": "https://jigyasa.example"}),
        ):
            with TestClient(backend_main.app) as client:
                response = client.get("/api/documents")
        self.assertEqual(response.status_code, 401)
        self.assertEqual(response.json()["detail"], "Sign in to access your study materials.")

    def test_production_rejects_state_changes_without_same_origin_header(self):
        with (
            patch.object(backend_main, "PRODUCTION", True),
            patch.object(backend_main, "_supabase_config", return_value=("https://example.supabase.co", "anon", "service")),
            patch.dict(os.environ, {"APP_ORIGIN": "https://jigyasa.example"}),
        ):
            with TestClient(backend_main.app) as client:
                response = client.post("/api/auth/logout")
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.json()["detail"], "Cross-origin requests are not allowed.")
        self.assertEqual(response.headers["x-content-type-options"], "nosniff")
        self.assertEqual(response.headers["cache-control"], "no-store")

    def test_hosted_document_queries_are_scoped_to_the_verified_user(self):
        user = {"id": "7c9c6f99-e916-40c6-9e67-f6813d553812"}
        with patch("backend.main._supabase_request", return_value=[]) as request:
            self.assertEqual(backend_main._user_documents(user), [])
        query = request.call_args.args[0]
        self.assertIn("owner_id=eq.7c9c6f99-e916-40c6-9e67-f6813d553812", query)
        self.assertNotIn("text_content", query)
        self.assertTrue(request.call_args.kwargs["service"])

    def test_hosted_chat_fetches_only_validated_selected_document_ids(self):
        user = {"id": "7c9c6f99-e916-40c6-9e67-f6813d553812"}
        document_id = "2e206f67-7fd8-4fab-a18a-944733c89e5b"
        with patch("backend.main._supabase_request", return_value=[]) as request:
            self.assertEqual(backend_main._user_documents(user, include_text=True, document_ids=[document_id]), [])
        query = request.call_args.args[0]
        self.assertIn("id=in.%282e206f67-7fd8-4fab-a18a-944733c89e5b%29", query)
        self.assertIn("text_content", query)

    def test_private_document_fields_are_never_returned_to_the_browser(self):
        document = {
            "id": "doc-id",
            "name": "notes.txt",
            "owner_id": "student-id",
            "text_content": "private extracted text",
            "storage_path": "student-id/doc-id.txt",
        }
        self.assertEqual(backend_main._public_document(document), {"id": "doc-id", "name": "notes.txt"})

    def test_daily_question_limit_is_enforced_by_the_database_counter(self):
        user = {"id": "7c9c6f99-e916-40c6-9e67-f6813d553812"}
        with patch.object(backend_main, "DAILY_AI_QUESTION_LIMIT", 20), patch(
            "backend.main._supabase_request", return_value=-20
        ) as request:
            with self.assertRaises(HTTPException) as context:
                backend_main._consume_daily_question(user)
        self.assertEqual(context.exception.status_code, 429)
        self.assertIn("20 AI questions", context.exception.detail)
        self.assertIn("consume_ai_question_quota", request.call_args.args[0])

    def test_failed_hosted_gemini_request_refunds_daily_question(self):
        user = {"id": "7c9c6f99-e916-40c6-9e67-f6813d553812"}
        document_id = "2e206f67-7fd8-4fab-a18a-944733c89e5b"
        documents = [{"id": document_id, "name": "notes.txt", "text_content": "A cell is a unit of life."}]
        with (
            patch.object(backend_main, "PRODUCTION", True),
            patch.dict(os.environ, {"GEMINI_API_KEY": "test-jigyasa-key"}),
            patch.object(backend_main, "_user_documents", return_value=documents),
            patch.object(backend_main, "_consume_daily_question") as consume,
            patch.object(backend_main, "_refund_daily_question") as refund,
            patch.object(backend_main, "_generate_gemini_answer", side_effect=HTTPException(503, "Gemini busy")),
        ):
            with self.assertRaises(HTTPException):
                backend_main.chat(
                    backend_main.ChatRequest(question="What is a cell?", document_ids=[document_id]),
                    user,
                )
        consume.assert_called_once_with(user)
        refund.assert_called_once_with(user)

    def test_hosted_session_cookies_are_http_only_and_secure(self):
        response = Response()
        session = {"access_token": "access-token", "refresh_token": "refresh-token", "expires_in": 3600}
        with patch.object(backend_main, "PRODUCTION", True):
            backend_main._set_session_cookies(response, session)
        cookies = response.headers.getlist("set-cookie")
        self.assertEqual(len(cookies), 2)
        self.assertTrue(all("HttpOnly" in cookie and "Secure" in cookie and "SameSite=lax" in cookie for cookie in cookies))

    def test_email_confirmation_verifies_tokens_before_setting_session_cookies(self):
        tokens = backend_main.AuthSessionTokens(access_token="a" * 32, refresh_token="r" * 32)
        response = Response()
        user = {"id": "7c9c6f99-e916-40c6-9e67-f6813d553812", "email": "student@example.test"}
        with (
            patch.object(backend_main, "PRODUCTION", True),
            patch("backend.main._supabase_request", return_value=user) as request,
        ):
            result = backend_main.confirm_email(tokens, response)
        self.assertEqual(result["id"], user["id"])
        self.assertEqual(request.call_args.args[:2], ("/auth/v1/user", "GET"))
        self.assertEqual(request.call_args.kwargs["access_token"], tokens.access_token)
        self.assertEqual(len(response.headers.getlist("set-cookie")), 2)

    def test_signup_accepts_immediate_supabase_session_response(self):
        response = Response()
        session = {
            "access_token": "a" * 32,
            "refresh_token": "r" * 32,
            "expires_in": 3600,
            "user": {"id": "7c9c6f99-e916-40c6-9e67-f6813d553812", "email": "student@example.test"},
        }
        credentials = backend_main.Credentials(email="student@example.test", password="safe-password-123")
        with patch.object(backend_main, "PRODUCTION", True), patch(
            "backend.main._auth_request", return_value=session
        ):
            result = backend_main.signup(credentials, response)
        self.assertTrue(result["authenticated"])
        self.assertFalse(result["email_confirmation_required"])
        self.assertEqual(result["id"], session["user"]["id"])
        self.assertEqual(len(response.headers.getlist("set-cookie")), 2)


class LocalApiTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        root = Path(self.temporary_directory.name)
        self.data_root_patch = patch.object(backend_main, "DATA_ROOT", root)
        self.document_root_patch = patch.object(backend_main, "DOCUMENT_ROOT", root / "documents")
        self.index_path_patch = patch.object(backend_main, "INDEX_PATH", root / "documents.json")
        self.data_root_patch.start()
        self.document_root_patch.start()
        self.index_path_patch.start()
        self.client = TestClient(backend_main.app)

    def tearDown(self):
        self.client.close()
        self.index_path_patch.stop()
        self.document_root_patch.stop()
        self.data_root_patch.stop()
        self.temporary_directory.cleanup()

    def test_upload_list_delete_document(self):
        upload = self.client.post(
            "/api/documents",
            files={"file": ("cell-notes.txt", b"Cells are the basic unit of life.", "text/plain")},
        )
        self.assertEqual(upload.status_code, 201)
        document = upload.json()
        self.assertEqual(document["name"], "cell-notes.txt")
        self.assertNotIn("text", document)

        listing = self.client.get("/api/documents")
        self.assertEqual([item["id"] for item in listing.json()], [document["id"]])

        deleted = self.client.delete(f"/api/documents/{document['id']}")
        self.assertEqual(deleted.status_code, 204)
        self.assertEqual(self.client.get("/api/documents").json(), [])

    def test_chat_without_gemini_key_is_explicitly_unavailable(self):
        upload = self.client.post(
            "/api/documents",
            files={"file": ("cell-notes.txt", b"Cells are the basic unit of life.", "text/plain")},
        )
        document_id = upload.json()["id"]
        with patch.dict(os.environ, {"GEMINI_API_KEY": ""}):
            answer = self.client.post(
                "/api/chat",
                json={"question": "What is a cell?", "document_ids": [document_id]},
            )
        self.assertEqual(answer.status_code, 503)
        self.assertIn("Gemini isn't connected", answer.json()["detail"])

    def test_chat_rejects_malformed_document_ids(self):
        answer = self.client.post(
            "/api/chat",
            json={"question": "What is a cell?", "document_ids": ["not-a-document-id"]},
        )
        self.assertEqual(answer.status_code, 422)

    def test_rejects_cross_origin_upload(self):
        response = self.client.post(
            "/api/documents",
            headers={"Origin": "https://another-site.example"},
            files={"file": ("cell-notes.txt", b"Cells are the basic unit of life.", "text/plain")},
        )
        self.assertEqual(response.status_code, 403)


if __name__ == "__main__":
    unittest.main()
